import io
import logging

import docx
import pypdfium2 as pdfium
from PIL import Image, ImageSequence, UnidentifiedImageError
from pypdf import PageObject, PdfReader
from pypdf.errors import PdfReadError

from app.core.config import Settings
from app.core.errors import ErrorCode, ProcessingError
from app.core.metrics import OCR_PAGES
from app.models.document import ExtractedText, SourceDocument, SourceFormat
from app.ocr.base import OcrError, OcrService
from app.services.file_type import detect_format

logger = logging.getLogger(__name__)


class TextExtractionService:
    """Obtiene el texto del documento. Usa la capa de texto cuando existe (PDF digital, DOCX, TXT)
    y recurre al OCR solo cuando hace falta (imágenes y páginas escaneadas)."""

    def __init__(self, ocr: OcrService, settings: Settings) -> None:
        self._ocr = ocr
        self._ocr_enabled = settings.ocr_enabled
        self._max_pages = settings.max_pages
        self._dpi = settings.ocr_dpi
        self._min_chars = settings.ocr_min_chars_per_page

    def extract(self, source: SourceDocument) -> ExtractedText:
        source_format = detect_format(source.content)
        match source_format:
            case SourceFormat.PDF:
                return self._from_pdf(source.content)
            case SourceFormat.IMAGE:
                return self._from_image(source.content)
            case SourceFormat.DOCX:
                return self._from_docx(source.content)
            case SourceFormat.TEXT:
                return self._from_text(source.content)

    # ---- PDF ----

    def _from_pdf(self, content: bytes) -> ExtractedText:
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted and not reader.decrypt(""):
                raise ProcessingError(ErrorCode.UNSUPPORTED_FORMAT, "Password-protected PDFs are not supported")
            total_pages = len(reader.pages)
        except PdfReadError as ex:
            raise ProcessingError(ErrorCode.CORRUPTED_DOCUMENT, "PDF file is corrupted or invalid") from ex

        warnings: list[str] = []
        pages = min(total_pages, self._max_pages)
        if total_pages > pages:
            warnings.append(f"Only the first {pages} of {total_pages} pages were processed")

        texts: list[str] = []
        ocr_pages = 0
        rendered: pdfium.PdfDocument | None = None
        try:
            for index in range(pages):
                try:
                    text = reader.pages[index].extract_text() or ""
                except Exception:  # pypdf puede fallar en páginas concretas con fuentes rotas
                    logger.warning("Text layer of page %d could not be read", index + 1)
                    text = ""
                # Solo se aplica OCR a páginas con poco texto que además contienen imágenes (escaneadas):
                # una página digital corta (una firma, un pie de página) no lo necesita
                if len(text.strip()) < self._min_chars and self._ocr_enabled and _has_images(reader.pages[index]):
                    rendered = rendered or pdfium.PdfDocument(content)
                    image = rendered[index].render(scale=self._dpi / 72).to_pil()
                    ocr_text = self._ocr_image(image)
                    ocr_pages += 1
                    # Si el OCR no encuentra más texto que la capa original, se conserva la original
                    if len(ocr_text.strip()) > len(text.strip()):
                        text = ocr_text
                texts.append(text)
        except pdfium.PdfiumError as ex:
            raise ProcessingError(ErrorCode.CORRUPTED_DOCUMENT, "PDF page could not be rendered for OCR") from ex
        finally:
            if rendered is not None:
                rendered.close()

        return ExtractedText("\n\n".join(texts), SourceFormat.PDF, total_pages, ocr_pages, warnings)

    # ---- Imágenes ----

    def _from_image(self, content: bytes) -> ExtractedText:
        if not self._ocr_enabled:
            raise ProcessingError(ErrorCode.OCR_UNAVAILABLE, "OCR is disabled and the document is an image")
        try:
            with Image.open(io.BytesIO(content)) as image:
                frames = [frame.copy() for frame in ImageSequence.Iterator(image)]
        except (UnidentifiedImageError, Image.DecompressionBombError, OSError) as ex:
            raise ProcessingError(ErrorCode.CORRUPTED_DOCUMENT, "Image file is corrupted or too large") from ex

        warnings: list[str] = []
        if len(frames) > self._max_pages:
            warnings.append(f"Only the first {self._max_pages} of {len(frames)} frames were processed")
        selected = frames[: self._max_pages]
        texts = [self._ocr_image(frame) for frame in selected]
        return ExtractedText("\n\n".join(texts), SourceFormat.IMAGE, len(frames), len(selected), warnings)

    # ---- DOCX y texto plano ----

    def _from_docx(self, content: bytes) -> ExtractedText:
        try:
            document = docx.Document(io.BytesIO(content))
        except Exception as ex:  # python-docx lanza distintos errores de zip/xml
            raise ProcessingError(ErrorCode.CORRUPTED_DOCUMENT, "DOCX file is corrupted or invalid") from ex
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return ExtractedText("\n".join(lines), SourceFormat.DOCX, page_count=1)

    def _from_text(self, content: bytes) -> ExtractedText:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252", errors="replace")
        return ExtractedText(text, SourceFormat.TEXT, page_count=1)

    def _ocr_image(self, image: Image.Image) -> str:
        if not self._ocr.is_available():
            raise ProcessingError(ErrorCode.OCR_UNAVAILABLE, f"OCR provider '{self._ocr.name}' is unavailable")
        try:
            text = self._ocr.extract_text(image)
        except OcrError as ex:
            raise ProcessingError(ErrorCode.OCR_UNAVAILABLE, str(ex)) from ex
        OCR_PAGES.inc()
        return text


def _has_images(page: PageObject) -> bool:
    """Indica si la página contiene imágenes (o formularios que pueden contenerlas), sin decodificarlas."""
    try:
        resources = page.get("/Resources")
        xobjects = resources.get_object().get("/XObject") if resources else None
        if not xobjects:
            return False
        xobjects = xobjects.get_object()
        return any(xobjects[name].get_object().get("/Subtype") in ("/Image", "/Form") for name in xobjects)
    except Exception:  # PDF con recursos mal formados: ante la duda se intenta el OCR
        return True
