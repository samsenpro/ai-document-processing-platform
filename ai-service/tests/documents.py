"""Genera archivos reales (PDF, DOCX, imágenes) en memoria para los tests."""

import io

import docx
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def text_pdf(pages: list[str]) -> bytes:
    """PDF digital: cada página tiene capa de texto."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for page in pages:
        y = 800
        for line in page.split("\n"):
            pdf.drawString(50, y, line)
            y -= 16
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def text_image(lines: list[str], fmt: str = "PNG") -> bytes:
    """Imagen con texto negro sobre blanco, legible para un OCR real."""
    font = ImageFont.load_default(size=40)
    image = Image.new("RGB", (1400, 90 + 70 * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(lines):
        draw.text((40, 40 + 70 * index), line, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, fmt)
    return buffer.getvalue()


def scanned_pdf(lines: list[str], pages: int = 1) -> bytes:
    """PDF sin capa de texto (solo imágenes), como el de un escáner."""
    frames = [Image.open(io.BytesIO(text_image(lines))).convert("RGB") for _ in range(pages)]
    buffer = io.BytesIO()
    frames[0].save(buffer, "PDF", save_all=True, append_images=frames[1:], resolution=150)
    return buffer.getvalue()


def docx_file(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
